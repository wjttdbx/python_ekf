"""Build formal patent specification and separate drawings with two embodiments."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


PATENT_DIR = Path(r"C:\Users\何逸凡\OneDrive\桌面\专利-EKF")
SPEC_SOURCE = PATENT_DIR / "100002说明书_插入仿真图_20260727.docx"
DRAWINGS_SOURCE = PATENT_DIR / "100003说明书附图.docx"
SPEC_OUTPUT = (
    PATENT_DIR / "100002说明书_双测量模式_估计距离判据_原生公式_20260727.docx"
)
DRAWINGS_OUTPUT = (
    PATENT_DIR / "100003说明书附图_双测量模式_估计距离判据_20260727.docx"
)

ROOT = Path(__file__).resolve().parents[1]
ANGLE_DIR = (
    ROOT / "outputs" / "EKF-SDRE-相对导航"
    / "仿真附图_仅测角_实施例参数_估计距离判据_20260727"
)
RANGE_DIR = (
    ROOT / "outputs" / "EKF-SDRE-相对导航"
    / "仿真附图_距离角度_实施例参数_估计距离判据_20260727"
)
DRAWING_IMAGES = [
    ANGLE_DIR / "图3_相对交会轨迹与距离.png",
    ANGLE_DIR / "图4_EKF状态估计误差.png",
    ANGLE_DIR / "图5_SDRE控制输入.png",
    RANGE_DIR / "图6_相对交会轨迹与距离.png",
    RANGE_DIR / "图7_EKF状态估计误差.png",
    RANGE_DIR / "图8_SDRE控制输入.png",
]


def _delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _set_body_format(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    for run in paragraph.runs:
        run.font.size = Pt(12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "宋体"
        )


def _append_body(document: Document, text: str, style):
    paragraph = document.add_paragraph(text, style=style)
    _set_body_format(paragraph)
    return paragraph


def _insert_body_after(anchor, text: str):
    paragraph = anchor._parent.add_paragraph()
    paragraph._p.getparent().remove(paragraph._p)
    anchor._p.addnext(paragraph._p)
    paragraph.style = anchor.style
    paragraph.add_run(text)
    _set_body_format(paragraph)
    return paragraph


def _insert_display_after(anchor, text: str):
    paragraph = _insert_body_after(anchor, text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_together = True
    return paragraph


def _prune_unused_document_images(path: Path) -> None:
    rels_name = "word/_rels/document.xml.rels"
    doc_name = "word/document.xml"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}

    document_xml = etree.fromstring(members[doc_name])
    used_ids = set(document_xml.xpath("//@r:embed", namespaces={"r": office_rel_ns}))
    relationships = etree.fromstring(members[rels_name])
    removed_targets: set[str] = set()
    for relationship in list(relationships):
        if (
            relationship.get("Type", "").endswith("/image")
            and relationship.get("Id") not in used_ids
        ):
            target = relationship.get("Target", "")
            removed_targets.add("word/" + target.lstrip("/"))
            relationships.remove(relationship)
    members[rels_name] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    temporary = path.with_suffix(".tmp.docx")
    with ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for name, content in members.items():
            if name not in removed_targets:
                target.writestr(name, content)
    temporary.replace(path)


def build_specification() -> None:
    document = Document(SPEC_SOURCE)

    # Formal CN patent specifications keep drawings in the separate drawings file.
    for paragraph in list(document.paragraphs):
        if paragraph._p.xpath(".//w:drawing"):
            _delete_paragraph(paragraph)
        elif paragraph.text.startswith(("图3  ", "图4  ", "图5  ")):
            _delete_paragraph(paragraph)
        elif paragraph.text.startswith("为检验初始估计偏差条件下的闭环运行能力"):
            _delete_paragraph(paragraph)

    figure_five = next(
        p for p in document.paragraphs
        if p.text.startswith("图5为本发明追踪航天器")
    )
    anchor = figure_five
    for text in (
        "图6为本发明距离-角度EKF-SDRE闭环相对交会轨迹与相对距离变化曲线。",
        "图7为本发明距离-角度EKF的位置和速度估计误差曲线。",
        "图8为本发明距离-角度测量模式下追踪航天器的SDRE控制加速度变化曲线。",
    ):
        paragraph = anchor._parent.add_paragraph()
        paragraph._p.getparent().remove(paragraph._p)
        anchor._p.addnext(paragraph._p)
        paragraph.style = anchor.style
        paragraph.add_run(text)
        _set_body_format(paragraph)
        anchor = paragraph

    initialization = next(
        p for p in document.paragraphs
        if p.text.startswith("设定追踪航天器和目标航天器的初始绝对状态")
    )
    init_detail = _insert_body_after(
        initialization,
        "本实施例将初始传感器误差对状态初始化的影响表示为初始状态估计误差"
        "δx0，并取x̂0=x0+δx0，其中δx0服从均值为零、协方差为"
        "P_EKF,0的高斯分布。采用随机数种子42得到本次确定性样本"
        "如下：",
    )
    error_vector = _insert_display_after(
        init_detail,
        "δx0=[0.036846,-0.125755,0.090745,0.00013133,-0.00027242,"
        "-0.00018182]",
    )
    seed_detail = _insert_body_after(
        error_vector,
        "其中前三项单位为km，后三项单位为km/s。初始真近点角取ν0=0 rad；"
        "后续测量噪声采用独立随机数种子43。",
    )
    _insert_body_after(
        seed_detail,
        "SDRE博弈调节因子取γ=√2。输入控制惩罚矩阵为R=10¹³I3，ARE求解"
        "和反馈控制实际采用的有效控制惩罚矩阵为"
        "Reff=R/(1-γ⁻²)=2×10¹³I3，相应地"
        "Reff⁻¹=(1-γ⁻²)R⁻¹=0.5×10⁻¹³I3。",
    )

    scale_paragraph = next(
        p for p in document.paragraphs
        if p.text.startswith("ARE数值尺度平衡预处理子步骤")
    )
    scale_paragraph.text = (
        "ARE数值尺度平衡预处理子步骤：本实施例的状态权重矩阵"
        "Qsdre=I6，其对角元素均为1；有效增益矩阵"
        "S=BReff⁻¹Bᵀ的非零对角元素为5×10⁻¹⁴量级。由于Qsdre与S的"
        "Frobenius范数相差约10¹³量级，直接构造哈密顿矩阵可能造成数值"
        "条件恶化，因此执行以下ARE数值尺度平衡预处理："
    )
    _set_body_format(scale_paragraph)

    are_solution = next(
        p for p in document.paragraphs
        if p.text.startswith("得到对称稳定矩阵P")
    )
    _insert_body_after(
        are_solution,
        "本实施例在上述ARE和控制律中以Reff替代输入参数R，即求解"
        "AᵀP+PA-PBReff⁻¹BᵀP+Qsdre=0，并计算"
        "u=-Reff⁻¹BᵀP x̂。目标航天器在状态传播中仍保持无控，"
        "其主动控制量恒为零。",
    )

    capture_step = next(
        p for p in document.paragraphs if p.text.startswith("计算当前相对距离")
    )
    capture_step.text = (
        "计算EKF当前相对位置估计值的范数作为星上交会判断距离d̂。若"
        "d̂<dmin（本实施例中dmin=100 m），则判定满足交会接近条件并结束"
        "闭环流程；否则返回步骤S2。仿真中同时记录真实相对距离d，用于离线"
        "评价估计距离判据的偏差，但真实相对距离不参与星上交会判断。"
    )
    _set_body_format(capture_step)

    technical_effect = next(
        p for p in document.paragraphs if p.text.startswith("本实施例技术效果：")
    )
    technical_effect.text = (
        "本实施例技术效果：按照上述参数并以EKF估计距离作为交会判断依据，"
        "仅测角EKF-SDRE闭环于6.392 h时首次满足d̂<100 m。该时刻估计距离"
        "为95.43 m，离线记录的真实距离为16.45 m，真实相对速度为"
        "0.1041 m/s；真实距离首次进入100 m范围的时刻为5.872 h。"
        "判定时刻的位置估计误差范数为78.98 m，速度估计误差范数为"
        "0.07764 m/s。控制加速度峰值为1.1693 m/s²，ARE求解过程中未发生"
        "回退。具体仿真曲线如图3至图5所示。"
    )
    _set_body_format(technical_effect)

    descriptions = {
        "图3示出了": (
            "图3示出了仅测角模式下的真实相对轨迹、EKF估计轨迹以及真实和"
            "估计相对距离随时间的变化。估计距离于6.392 h时首次小于"
            "0.1 km；同一时刻真实距离为16.45 m。图中的判定时刻按照估计"
            "距离确定。"
        ),
        "图4示出了": (
            "图4示出了仅测角EKF的位置和速度估计误差范数及相应协方差尺度。"
            "EKF初始估计误差由P_EKF,0和随机数种子42生成；随着含噪量测进入"
            "滤波过程，估计误差发生变化。判定时刻的位置估计误差为78.98 m，"
            "速度估计误差为0.07764 m/s。"
        ),
        "图5示出了": (
            "图5示出了仅测角模式下三个LVLH方向的控制加速度及其合成范数。"
            "本实施例未施加控制幅值截断，控制加速度峰值为1.1693 m/s²，"
            "控制加速度范数全程平均值为0.09891 m/s²。"
        ),
    }
    for prefix, text in descriptions.items():
        paragraph = next(
            p for p in document.paragraphs if p.text.startswith(prefix)
        )
        paragraph.text = text
        _set_body_format(paragraph)

    embodiment_index, embodiment_two = next(
        (i, p) for i, p in enumerate(document.paragraphs)
        if p.text.startswith("实施例二：距离+角度测量模式")
    )
    body_style = document.paragraphs[embodiment_index + 1].style
    for paragraph in list(document.paragraphs)[embodiment_index + 1:]:
        _delete_paragraph(paragraph)

    paragraphs = (
        "本实施例采用与实施例一相同的地球椭圆轨道交会场景、真实初始相对状态、"
        "EKF初始协方差、由P_EKF,0生成的初始状态估计误差、过程噪声、"
        "控制权重、控制步长、"
        "交会距离阈值、初始真近点角ν0=0 rad、博弈调节因子γ=√2以及随机数"
        "种子配置。初始误差随机数种子为42，后续测量噪声随机数种子为43。"
        "区别在于追踪航天器在测角传感器之外还配备测距传感器，以同时获得"
        "相对距离、方位角和俯仰角测量值。",
        "本实施例的测距噪声标准差为0.01 km，即10 m；方位角和俯仰角噪声"
        "标准差均为0.008°，即约1.4×10⁻⁴ rad。因此，距离-角度量测噪声"
        "协方差矩阵取Rz=diag(0.01²,(1.4×10⁻⁴)²,(1.4×10⁻⁴)²)，"
        "其中距离分量单位为km²，两个角度分量单位为rad²。",
        "在步骤S7中，量测向量由实施例一的二维方位角-俯仰角量测扩展为"
        "三维距离-方位角-俯仰角量测。量测函数依次输出相对距离ρ、方位角az"
        "和俯仰角el；相应量测雅可比矩阵增加距离对三个相对位置分量的偏导数"
        "[Δx/ρ, Δy/ρ, Δz/ρ]。EKF根据量测噪声协方差矩阵和量测向量的维度"
        "选择对应的量测函数及雅可比矩阵，预测、更新和SDRE控制步骤无需改变。",
        "按照上述参数并以EKF估计距离作为交会判断依据，距离-角度EKF-SDRE"
        "闭环于8.281 h时首次满足d̂<100 m。该时刻估计距离为93.47 m，"
        "离线记录的真实距离为109.11 m，真实相对速度为0.07295 m/s；位置"
        "估计误差范数为15.64 m，速度估计误差范数为0.1588 m/s。控制加速度"
        "峰值为1.1693 m/s²，控制加速度范数全程平均值为0.07481 m/s²，"
        "ARE求解过程中未发生回退。",
        "图6示出了距离-角度测量模式下的真实相对轨迹、EKF估计轨迹以及相对"
        "距离随时间的变化。估计距离于8.281 h时首次达到0.1 km交会判断"
        "阈值；同一时刻真实距离为109.11 m。因此，该判据在本次噪声序列下"
        "产生约9.11 m的提前判定偏差。",
        "图7示出了距离-角度EKF的位置和速度估计误差范数及相应协方差尺度。"
        "在测距信息直接约束径向位置后，交会判定时刻的位置估计误差为"
        "15.64 m，速度估计误差为0.1588 m/s。",
        "图8示出了距离-角度测量模式下三个LVLH方向的控制加速度及其合成范数。"
        "结合图6至图8可知，本发明在不改变EKF-SDRE闭环主体逻辑的条件下，"
        "能够适配距离-角度三维量测并按照估计距离完成交会判断。真实距离与"
        "估计距离在阈值附近存在偏差，工程实施时可结合安全裕量或连续多周期"
        "确认策略进行处理；上述结果仅对应本实施例给定参数和随机数种子。",
    )
    for text in paragraphs:
        paragraph = _append_body(document, text, body_style)
        if text.startswith("本实施例的测距噪声标准差"):
            paragraph.paragraph_format.keep_together = True

    document.core_properties.title = (
        "一种基于EKF-SDRE紧耦合的航天器相对导航控制方法及系统"
    )
    document.save(SPEC_OUTPUT)
    _prune_unused_document_images(SPEC_OUTPUT)


def build_drawings() -> None:
    for image in DRAWING_IMAGES:
        if not image.exists():
            raise FileNotFoundError(image)

    document = Document(DRAWINGS_SOURCE)
    # The source's second page is template guidance, not a patent drawing.
    for paragraph in list(document.paragraphs)[7:]:
        _delete_paragraph(paragraph)

    section = document.sections[0]
    section.header.paragraphs[0].text = "说明书附图"
    section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    usable_width = min(
        Cm(16.0),
        section.page_width - section.left_margin - section.right_margin,
    )
    for number, image_path in enumerate(DRAWING_IMAGES, start=3):
        page_break = document.add_paragraph()
        page_break.add_run().add_break(WD_BREAK.PAGE)

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(0)
        image_paragraph.paragraph_format.space_after = Pt(4)
        image_paragraph.paragraph_format.keep_with_next = True
        shape = image_paragraph.add_run().add_picture(
            str(image_path), width=usable_width
        )
        shape._inline.docPr.set("title", f"图{number}")
        shape._inline.docPr.set("descr", image_path.stem)

        caption = document.add_paragraph(f"图{number}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(0)
        caption.paragraph_format.first_line_indent = Cm(0)
        for run in caption.runs:
            run.font.size = Pt(10.5)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), "宋体"
            )

    document.core_properties.title = "说明书附图"
    document.save(DRAWINGS_OUTPUT)


def main() -> None:
    build_specification()
    build_drawings()
    print(SPEC_OUTPUT)
    print(DRAWINGS_OUTPUT)


if __name__ == "__main__":
    main()
