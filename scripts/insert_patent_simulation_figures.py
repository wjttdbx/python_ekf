"""Insert verified simulation figures into the EKF-SDRE patent specification."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"C:\Users\何逸凡\OneDrive\桌面\专利-EKF\100002说明书.docx")
OUTPUT = Path(
    r"C:\Users\何逸凡\OneDrive\桌面\专利-EKF"
    r"\100002说明书_插入仿真图_20260727.docx"
)
FIGURE_DIR = (
    Path(__file__).resolve().parents[1]
    / "outputs"
    / "EKF-SDRE-相对导航"
    / "仿真附图_原始Pekf_固定扰动_20260727"
)

FIGURES = (
    (
        FIGURE_DIR / "图3_相对交会轨迹与距离.png",
        "图3  仅测角EKF-SDRE闭环相对交会轨迹与相对距离变化曲线",
        (
            "图3示出了本实施例中追踪航天器在LVLH坐标系下的相对运动轨迹"
            "以及相对距离随时间的变化曲线。初始相对距离约为866 km。随着"
            "EKF-SDRE闭环控制过程推进，真实相对距离总体减小，并于5.872 h时"
            "首次小于0.1 km的交会距离阈值；首次进入该阈值时的真实相对速度"
            "为0.2398 m/s。"
        ),
        "仅测角EKF-SDRE闭环相对交会轨迹与距离曲线",
    ),
    (
        FIGURE_DIR / "图4_EKF状态估计误差.png",
        "图4  仅测角EKF的位置和速度估计误差曲线",
        (
            "图4示出了仅测角EKF的位置估计误差范数、速度估计误差范数及相应"
            "协方差尺度曲线。仿真初始位置估计扰动范数为7.681 km，初始速度"
            "估计扰动范数为1.500 m/s；首次进入交会距离阈值时，位置估计误差"
            "范数为0.299 km，速度估计误差范数为0.516 m/s。误差在交会过程中"
            "存在阶段性波动，但能够为SDRE控制器持续提供闭环状态估计。"
        ),
        "仅测角EKF位置和速度估计误差曲线",
    ),
    (
        FIGURE_DIR / "图5_SDRE控制输入.png",
        "图5  追踪航天器的SDRE控制加速度变化曲线",
        (
            "图5示出了追踪航天器在三个LVLH方向上的控制加速度及其合成范数。"
            "本实施例未对控制量施加幅值截断，仿真得到的控制加速度峰值为"
            "1.1895 m/s²，控制加速度范数的全程平均值为0.1076 m/s²。结合图3"
            "至图5可知，在仅使用方位角和俯仰角测量的条件下，本实施例形成了"
            "状态估计、SDRE反馈控制与非线性轨道传播的闭环运行过程。"
        ),
        "追踪航天器SDRE控制加速度曲线",
    ),
)


def _find_paragraph(document: Document, predicate) -> object:
    for paragraph in document.paragraphs:
        if predicate(paragraph.text):
            return paragraph
    raise ValueError("未找到目标段落。")


def _set_body_format(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    for run in paragraph.runs:
        run.font.size = Pt(12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "宋体"
        )


def _add_body_before(anchor, text: str):
    paragraph = anchor.insert_paragraph_before(text, style=anchor.style)
    _set_body_format(paragraph)
    return paragraph


def _add_figure_before(anchor, image_path: Path, caption: str, description: str,
                       alt_text: str, width_cm: float, page_break: bool = False) -> None:
    image_paragraph = anchor.insert_paragraph_before(style=anchor.style)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.paragraph_format.page_break_before = page_break
    run = image_paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Cm(width_cm))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", caption)

    caption_paragraph = anchor.insert_paragraph_before(caption, style=anchor.style)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(3)
    caption_paragraph.paragraph_format.first_line_indent = Cm(0)
    for caption_run in caption_paragraph.runs:
        caption_run.font.size = Pt(10.5)
        caption_run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "宋体"
        )

    _add_body_before(anchor, description)


def _replace_p0_diagonal(paragraph) -> None:
    math_nodes = paragraph._p.xpath(".//m:oMath")
    target = None
    for math_node in math_nodes:
        texts = math_node.xpath(
            ".//*[local-name()='t']/text()"
        )
        if "diag" in texts and texts.count("100") == 3 and texts.count("1") >= 3:
            target = math_node
            break
    if target is None:
        raise ValueError("未找到原P0=diag(100,100,100,1,1,1)公式。")

    value_nodes = target.xpath(
        ".//*[local-name()='d']/*[local-name()='e']"
        "/*[local-name()='r']/*[local-name()='t']"
    )
    numeric_nodes = [
        node for node in value_nodes
        if (node.text or "").strip() not in {",", ""}
    ]
    old_values = [node for node in numeric_nodes if node.text in {"100", "1"}]
    if len(old_values) != 6:
        raise ValueError(f"P0公式元素数量异常：{len(old_values)}")
    replacements = ["0.0146216"] * 3 + ["1.94955×10⁻⁸"] * 3
    for node, replacement in zip(old_values, replacements):
        node.text = replacement

    for text_node in paragraph._p.xpath(".//w:t"):
        if text_node.text and "反映初始相对状态先验信息较弱" in text_node.text:
            text_node.text = text_node.text.replace(
                "反映初始相对状态先验信息较弱",
                "由初始距离与测角噪声尺度确定",
            )


def main() -> None:
    for image_path, *_ in FIGURES:
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    document = Document(SOURCE)

    appendix_blank = document.paragraphs[49]
    if appendix_blank.text:
        raise ValueError("附图说明后的定位段落发生变化。")
    for text in (
        "图3为本发明仅测角EKF-SDRE闭环相对交会轨迹与相对距离变化曲线。",
        "图4为本发明仅测角EKF的位置和速度估计误差曲线。",
        "图5为本发明追踪航天器的SDRE控制加速度变化曲线。",
    ):
        paragraph = appendix_blank.insert_paragraph_before(text, style=appendix_blank.style)
        _set_body_format(paragraph)

    initialization = _find_paragraph(
        document,
        lambda text: text.startswith("设定追踪航天器和目标航天器的初始绝对状态"),
    )
    _replace_p0_diagonal(initialization)
    perturbation_paragraph = initialization._parent.add_paragraph()
    perturbation_paragraph._p.getparent().remove(perturbation_paragraph._p)
    initialization._p.addnext(perturbation_paragraph._p)
    perturbation_paragraph.style = initialization.style
    perturbation_paragraph.add_run(
        "为检验初始估计偏差条件下的闭环运行能力，本实施例在真实初始相对状态"
        "基础上加入固定估计扰动：位置扰动为[5, -5, 3] km，速度扰动为"
        "[0.001, -0.001, 0.0005] km/s。"
    )
    _set_body_format(perturbation_paragraph)

    technical_effect = _find_paragraph(
        document,
        lambda text: text.startswith("本实施例技术效果："),
    )
    technical_effect.text = (
        "本实施例技术效果：在上述参数配置及固定初始估计扰动下，追踪航天器"
        "利用仅测角EKF-SDRE闭环于5.872 h时首次进入100 m交会距离范围，"
        "进入时真实相对速度为0.2398 m/s。控制加速度峰值为1.1895 m/s²，"
        "首次进入交会范围时的位置估计误差范数为0.299 km。ARE求解过程中"
        "未发生回退。具体仿真曲线如图3至图5所示。"
    )
    _set_body_format(technical_effect)

    embodiment_two = _find_paragraph(
        document,
        lambda text: text.startswith("实施例二：距离+角度测量模式"),
    )
    usable_width_cm = min(
        16.0,
        (document.sections[0].page_width
         - document.sections[0].left_margin
         - document.sections[0].right_margin) / 360000.0,
    )
    for index, (image_path, caption, description, alt_text) in enumerate(FIGURES):
        _add_figure_before(
            embodiment_two,
            image_path,
            caption,
            description,
            alt_text,
            width_cm=usable_width_cm,
            page_break=(index == 0),
        )

    document.core_properties.title = "一种基于EKF-SDRE紧耦合的航天器相对导航控制方法及系统"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
